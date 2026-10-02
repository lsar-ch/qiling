"""本地办公执行器：读取常见办公文件并返回内容摘要。

B 为主、A 兜底：外部 Agent（Codex/WorkBuddy/其它 ACP 执行器）不可用时，用本地库
处理窄场景办公任务（读取内容、数据摘要、简单统计）。覆盖范围：

- Office：docx / xlsx / pptx（读写由外部 Agent 负责，这里只读）
- PDF：文本与元数据提取（扫描件需 OCR，本地不做）
- 图片：尺寸/格式/EXIF 元数据（看图内容需外部模型）
- 邮件：.eml 的头/正文/附件清单

写作：没有源文件（或任务本身就是要新建）时，调用已配置的模型写 Markdown，
再本地渲染成 docx / xlsx / pptx / pdf / md —— 不依赖外部 Agent 也能产出文件。
复杂编辑既有文件、跨格式转换仍走外部 Agent。
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from deskmate_contracts.agent import AgentInfo

from app.agents.base import (
    AgentAdapter,
    AgentTaskRequest,
    AgentTaskResult,
)
from app.agents.office_writer import (
    default_filename,
    parse_markdown,
    title_of,
    write_file,
)
from app.llm.base import LLMProvider

logger = logging.getLogger(__name__)

_SUFFIXES = {
    "word": {".docx"},
    "excel": {".xlsx", ".xlsm"},
    "ppt": {".pptx"},
    "pdf": {".pdf"},
    "image": {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"},
    "email": {".eml"},
}

_MAX_PDF_PAGES = 20
_MAX_BODY_CHARS = 4000

# 明确要「新建」而不是「读现有的」时，即使目录里有同类型文件也走写作分支。
_CREATE_HINTS = re.compile(r"写一份|写一篇|起草|生成一份|生成一份|新建|创建|做一份|做一份|拟一份|帮我写")

_AUTHOR_SYSTEM = (
    "你是 DeskMate 的办公写作助手。只输出 Markdown 正文，不要寒暄、不要解释你的做法。"
    "标题用 # / ## / ###，要点用 - 开头，表格用 | 列 | 列 | 的 Markdown 表格。"
    "内容要具体、可直接交给上级或客户，不要写占位符。"
)


def _find_office_files(cwd: str, work_kind: str, limit: int = 5) -> list[Path]:
    """在工作目录下找到匹配 work_kind 的办公文件，按修改时间倒序。"""
    root = Path(cwd)
    if not root.is_dir():
        return []
    suffixes = _SUFFIXES.get(work_kind, set())
    found: list[tuple[float, Path]] = []
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in suffixes:
            if any(part.startswith(".") for part in path.parts):
                continue
            if path.name.startswith("~"):
                continue
            try:
                found.append((path.stat().st_mtime, path))
            except OSError:
                continue
    found.sort(reverse=True)
    return [path for _, path in found[:limit]]


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs[:80]:
        text = para.text.strip()
        if text:
            style = para.style.name if para.style else ""
            if "Heading" in style:
                lines.append(f"## {text}")
            else:
                lines.append(text)
    for i, table in enumerate(doc.tables[:5]):
        lines.append(f"\n[表格 {i + 1}]")
        for row in table.rows[:20]:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
    if not lines:
        return "（文档为空或无可提取文本）"
    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for name in wb.sheetnames[:10]:
        ws = wb[name]
        parts.append(f"## 工作表：{name}")
        rows = list(ws.iter_rows(max_row=30, values_only=True))
        if not rows:
            parts.append("（空表）")
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        parts.append(" | ".join(header))
        parts.append("-" * 40)
        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            parts.append(" | ".join(cells))
        total_rows = ws.max_row or len(rows)
        if total_rows > 30:
            parts.append(f"... 共 {total_rows} 行，仅显示前 30 行")
    wb.close()
    return "\n".join(parts) if parts else "（工作簿为空）"


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(list(prs.slides)[:30]):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            parts.append(f"## 第 {i + 1} 页")
            parts.extend(texts)
    if not parts:
        return "（演示文稿为空或无可提取文本）"
    return "\n".join(parts)


def _read_pdf(path: Path, max_pages: int = _MAX_PDF_PAGES) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    total = len(reader.pages)
    parts: list[str] = []
    for index in range(min(total, max_pages)):
        try:
            text = (reader.pages[index].extract_text() or "").strip()
        except Exception as exc:  # 单页损坏不应让整份 PDF 失败
            logger.warning("PDF 第 %s 页解析失败（%s）：%s", index + 1, path.name, exc)
            text = ""
        if text:
            parts.append(f"## 第 {index + 1} 页\n{text}")
    if not parts:
        return "（PDF 无可提取文本；可能是扫描件，需要 OCR）"
    if total > max_pages:
        parts.append(f"... 共 {total} 页，仅提取前 {max_pages} 页")
    return "\n\n".join(parts)


def _read_image(path: Path) -> str:
    from PIL import ExifTags, Image

    lines: list[str] = []
    with Image.open(str(path)) as img:
        lines.append(f"- 格式：{img.format or path.suffix.upper().lstrip('.')}")
        lines.append(f"- 尺寸：{img.width} × {img.height} 像素")
        lines.append(f"- 色彩模式：{img.mode}")
        try:
            exif = img.getexif()
        except Exception:
            exif = None
        if exif:
            wanted = {
                "Make", "Model", "DateTime", "DateTimeOriginal",
                "Software", "Orientation", "FNumber", "ExposureTime", "ISOSpeedRatings",
            }
            for tag_id, value in exif.items():
                name = ExifTags.TAGS.get(tag_id, str(tag_id))
                if name in wanted and value not in (None, ""):
                    lines.append(f"- {name}：{value}")
    lines.append("- 备注：本地只读元数据；识别图片里写了什么需要 OCR 或云端模型。")
    return "## 图片信息\n" + "\n".join(lines)


def _read_email(path: Path, max_chars: int = _MAX_BODY_CHARS) -> str:
    import email
    from email import policy

    with path.open("rb") as handle:
        msg = email.message_from_binary_file(handle, policy=policy.default)

    parts: list[str] = ["## 邮件"]
    for label, key in (
        ("主题", "Subject"),
        ("发件人", "From"),
        ("收件人", "To"),
        ("抄送", "Cc"),
        ("日期", "Date"),
    ):
        value = msg.get(key)
        if value:
            parts.append(f"- {label}：{value}")

    body = _email_body_text(msg)
    if body:
        if len(body) > max_chars:
            body = body[:max_chars] + "\n...（正文过长，已截断）"
        parts.append("\n## 正文\n" + body)

    attachments = [
        part.get_filename()
        for part in msg.walk()
        if part.get_content_disposition() == "attachment" and part.get_filename()
    ]
    if attachments:
        parts.append("\n## 附件\n" + "\n".join(f"- {name}" for name in attachments))
    return "\n".join(parts)


def _email_body_text(msg) -> str:
    """优先 text/plain；没有就取 text/html 去标签。附件部分跳过。"""
    if not msg.is_multipart():
        try:
            content = msg.get_content()
        except Exception:
            return ""
        text = str(content or "")
        return _strip_html(text) if msg.get_content_type() == "text/html" else text.strip()
    for part in msg.walk():
        if part.get_content_disposition() == "attachment":
            continue
        if part.get_content_type() != "text/plain":
            continue
        try:
            content = part.get_content()
        except Exception:
            continue
        if content and str(content).strip():
            return str(content).strip()
    for part in msg.walk():
        if part.get_content_disposition() == "attachment":
            continue
        if part.get_content_type() == "text/html":
            try:
                return _strip_html(str(part.get_content() or "")).strip()
            except Exception:
                continue
    return ""


def _strip_html(text: str) -> str:
    return re.sub(r"<[^>]+>", " ", text)


_READERS = {
    "word": _read_docx,
    "excel": _read_xlsx,
    "ppt": _read_pptx,
    "pdf": _read_pdf,
    "image": _read_image,
    "email": _read_email,
}


class LocalOfficeAdapter(AgentAdapter):
    """本地办公执行器：读常见办公文件，或调用已配置模型写作并落盘。

    provider 为 None 时只具备读取能力（写作分支会给出明确提示，不静默失败）。
    """

    name = "local-office"
    adapter_id = "local-office"
    display_name = "本地办公"
    executable = True
    default_eligible = False

    def __init__(self, provider: LLMProvider | None = None) -> None:
        self.provider = provider

    async def detect(self) -> AgentInfo:
        return AgentInfo(
            name=self.name,
            label=self.display_name,
            adapter=self.adapter_id,
            integration_status="ready",
            available=True,
            enabled=True,
            default_eligible=self.default_eligible,
            protocol="builtin",
            capabilities=[
                "office",
                "read",
                "summarize",
                "write",
                "pdf",
                "image",
                "email",
            ],
            tags=["builtin", "local"],
            detail=(
                "本地执行器：docx/xlsx/pptx/PDF 文本提取、图片元数据、.eml 邮件解析；"
                "并在已配置模型时可直接写作产出 docx/xlsx/pptx/pdf"
                + ("" if self.provider is not None else "（当前未配置模型，仅可读取）")
            ),
        )

    async def run(
        self,
        request: AgentTaskRequest,
        on_update,
    ) -> AgentTaskResult:
        work_kind = request.config.get("work_kind", "")
        if work_kind not in _READERS:
            return AgentTaskResult(
                status="failed",
                error=f"本地执行器不支持类型：{work_kind}",
            )
        files = _find_office_files(request.cwd, work_kind)
        wants_create = bool(_CREATE_HINTS.search(request.prompt or ""))
        if not files or wants_create:
            return await self._author(request, work_kind, on_update)
        reader = _READERS[work_kind]
        summaries: list[str] = []
        for path in files[:3]:
            try:
                await on_update(
                    "progress",
                    {"message": f"正在读取 {path.name}..."},
                )
                content = reader(path)
                if len(content) > 2000:
                    content = content[:2000] + "\n...（内容过长，已截断）"
                summaries.append(f"### {path.name}\n\n{content}")
            except Exception as exc:
                logger.warning("读取 %s 失败: %s", path, exc)
                summaries.append(f"### {path.name}\n\n读取失败：{exc}")
        if not summaries:
            return AgentTaskResult(
                status="failed",
                error="所有文件读取均失败",
            )
        return AgentTaskResult(
            status="completed",
            message="\n\n".join(summaries),
        )

    # ------------------------------------------------------------------ 写作

    async def _author(
        self,
        request: AgentTaskRequest,
        work_kind: str,
        on_update,
    ) -> AgentTaskResult:
        """调用已配置模型写内容，再本地渲染成目标格式。"""
        if self.provider is None:
            return AgentTaskResult(
                status="failed",
                error=(
                    f"本地未找到 {_SUFFIXES.get(work_kind)} 源文件，且未配置模型无法写作。"
                    "请到「设置 → 模型」配置文本模型，或改用外部执行器。"
                ),
            )
        await on_update("progress", {"message": "正在用模型起草内容..."})
        try:
            result = await self.provider.chat(
                [
                    {"role": "system", "content": _AUTHOR_SYSTEM},
                    {"role": "user", "content": request.prompt},
                ]
            )
        except Exception as exc:
            logger.warning("本地写作调用模型失败: %s", exc)
            return AgentTaskResult(status="failed", error=f"调用模型失败：{exc}")

        text = (getattr(result, "text", "") or "").strip()
        if not text:
            return AgentTaskResult(status="failed", error="模型没有返回可用内容")
        blocks = parse_markdown(text)
        title = title_of(blocks, fallback=f"DeskMate {work_kind}")
        target = Path(request.cwd) / default_filename(work_kind, title)
        try:
            write_file(target, work_kind, blocks, raw_text=text)
        except Exception as exc:
            logger.warning("写文件失败 %s: %s", target, exc)
            return AgentTaskResult(status="failed", error=f"写入文件失败：{exc}")
        await on_update("progress", {"message": f"已生成 {target.name}"})
        return AgentTaskResult(
            status="completed",
            message=f"已生成文件：{target}\n\n标题：{title}\n\n{text[:1500]}",
        )
