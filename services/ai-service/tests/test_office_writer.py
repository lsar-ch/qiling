"""本地写作：Markdown 解析 + 各格式渲染 + 本地执行器写作分支。"""

from __future__ import annotations

import asyncio
from pathlib import Path

from app.agents import office_writer
from app.agents.base import AgentTaskRequest
from app.agents.local_office import LocalOfficeAdapter
from app.agents.office_writer import (
    default_filename,
    docx_template_context,
    parse_markdown,
    title_of,
    typst_source,
    write_file,
)
from app.llm.base import ChatResult

SAMPLE = """# 季度总结

本季度完成了三件事。

## 关键指标

- 交付 3 个模块
- 缺陷率下降 20%

## 数据

| 月份 | 收入 |
| --- | --- |
| 7月 | 100 |
| 8月 | 120 |
"""


class FakeProvider:
    """最小 provider：只回固定 Markdown。"""

    def __init__(self, text: str = SAMPLE) -> None:
        self.text = text
        self.calls: list[list[dict]] = []

    @property
    def display_name(self) -> str:
        return "fake"

    async def chat(self, messages, *, tools=None):
        self.calls.append(messages)
        return ChatResult(text=self.text)

    async def stream_chat(self, messages, *, temperature=None, max_tokens=None):
        yield self.text


class TestParser:
    def test_parses_headings_bullets_table(self):
        kinds = [block.kind for block in parse_markdown(SAMPLE)]
        assert kinds == ["h1", "paragraph", "h2", "bullet", "bullet", "h2", "table"]

    def test_table_separator_row_is_dropped(self):
        table = next(b for b in parse_markdown(SAMPLE) if b.kind == "table")
        assert table.rows[0] == ["月份", "收入"]
        assert len(table.rows) == 3

    def test_title_falls_back_when_empty(self):
        assert title_of(parse_markdown(SAMPLE)) == "季度总结"
        assert title_of([], fallback="兜底标题") == "兜底标题"

    def test_default_filename_strips_path_characters(self):
        assert default_filename("word", "../../etc/passwd") .endswith(".docx")
        assert "/" not in default_filename("word", "../evil/name")


class TestRenderers:
    def test_writes_docx(self, tmp_path):
        target = tmp_path / "out.docx"
        write_file(target, "word", parse_markdown(SAMPLE), SAMPLE)
        from docx import Document

        texts = [p.text for p in Document(str(target)).paragraphs]
        assert "季度总结" in texts
        assert "交付 3 个模块" in texts

    def test_writes_xlsx_with_table_sheet(self, tmp_path):
        target = tmp_path / "out.xlsx"
        write_file(target, "excel", parse_markdown(SAMPLE), SAMPLE)
        from openpyxl import load_workbook

        wb = load_workbook(str(target))
        assert "表1" in wb.sheetnames
        assert wb["表1"]["A1"].value == "月份"

    def test_writes_pptx(self, tmp_path):
        target = tmp_path / "out.pptx"
        write_file(target, "ppt", parse_markdown(SAMPLE), SAMPLE)
        from pptx import Presentation

        prs = Presentation(str(target))
        titles = [s.shapes.title.text for s in prs.slides]
        assert "季度总结" in titles

    def test_writes_pdf(self, tmp_path):
        target = tmp_path / "out.pdf"
        write_file(target, "pdf", parse_markdown(SAMPLE), SAMPLE)
        assert target.read_bytes()[:5] == b"%PDF-"

    def test_writes_markdown_for_unknown_kind(self, tmp_path):
        target = tmp_path / "out.md"
        write_file(target, "general", parse_markdown(SAMPLE), SAMPLE)
        assert "季度总结" in target.read_text(encoding="utf-8")


class TestOptionalEngines:
    """装了 typst / docxtpl 才启用的增强链路；没装必须安静回落内置渲染器。"""

    def test_typst_source_escapes_markup_and_keeps_bullets_together(self) -> None:
        source = typst_source(parse_markdown("# 标题\n\n- 一\n- 二\n\n用 # 和 * 号强调\n"))
        assert "= 标题" in source
        # 连续列表项之间不能有空行，否则 Typst 会拆成多个列表、行距变松
        assert "- 一\n- 二" in source
        assert "\\#" in source and "\\*" in source

    def test_typst_source_renders_table_with_header(self) -> None:
        source = typst_source(parse_markdown("| 月份 | 收入 |\n| --- | --- |\n| 7月 | 100 |\n"))
        assert "columns: 2" in source
        assert "table.header([月份], [收入])" in source
        assert "[7月], [100]," in source

    def test_docx_template_context_groups_sections(self) -> None:
        context = docx_template_context(
            parse_markdown("# 标题\n\n前言\n\n## 段一\n\n- 要点\n")
        )
        assert context["title"] == "标题"
        assert [section["heading"] for section in context["sections"]] == ["", "段一"]
        assert context["sections"][0]["paragraphs"] == ["前言"]
        assert context["sections"][1]["bullets"] == ["要点"]

    def test_engines_stand_down_when_not_installed(self, tmp_path, monkeypatch) -> None:
        blocks = parse_markdown("# 标题\n")
        monkeypatch.setattr(office_writer.shutil, "which", lambda _name: None)
        monkeypatch.setenv("DESKMATE_OFFICE_TEMPLATES", str(tmp_path / "no-such-dir"))
        assert office_writer._write_pdf_via_typst(tmp_path / "x.pdf", blocks) is False
        assert office_writer._write_docx_via_template(tmp_path / "x.docx", blocks) is False

    def test_write_file_still_produces_files_without_engines(self, tmp_path) -> None:
        blocks = parse_markdown("# 标题\n\n正文一段。\n")
        docx = write_file(tmp_path / "a.docx", "word", blocks)
        pdf = write_file(tmp_path / "a.pdf", "pdf", blocks)
        assert docx.stat().st_size > 0
        assert pdf.read_bytes()[:5] == b"%PDF-"


class TestAuthoringBranch:
    def test_local_adapter_writes_docx_when_no_source_file(self, tmp_path):
        provider = FakeProvider()
        adapter = LocalOfficeAdapter(provider=provider)

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="t1",
                    prompt="帮我写一份季度总结的 Word 文档",
                    cwd=str(tmp_path),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        produced = list(tmp_path.glob("*.docx"))
        assert len(produced) == 1
        assert provider.calls, "应当调用模型"

    def test_local_adapter_without_provider_fails_clearly(self, tmp_path):
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="t2",
                    prompt="帮我写一份报告",
                    cwd=str(tmp_path),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "failed"
        assert "未配置模型" in result.error

    def test_reading_still_wins_when_source_file_exists(self, tmp_path):
        from docx import Document

        Document().save(str(tmp_path / "src.docx"))
        doc = Document(str(tmp_path / "src.docx"))
        doc.add_paragraph("源文件正文")
        doc.save(str(tmp_path / "src.docx"))

        provider = FakeProvider()
        adapter = LocalOfficeAdapter(provider=provider)

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="t3",
                    prompt="总结这个文档",
                    cwd=str(tmp_path),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "源文件正文" in result.message
        assert not provider.calls, "读现有文件时不该调用模型"


class TestLayoutStandard:
    """standard_business_brief 的排版 token 必须真的落进产物。

    `TestRenderers` 只证明"能打开"；这里断言页面/字号/颜色/编号/表格几何/页脚域，
    否则标准只是 docstring 里的一段话。
    """

    def _docx(self, tmp_path):
        target = tmp_path / "layout.docx"
        write_file(target, "word", parse_markdown(SAMPLE), SAMPLE)
        from docx import Document

        return Document(str(target))

    def test_docx_page_setup(self, tmp_path):
        section = self._docx(tmp_path).sections[0]
        assert round(section.page_width.inches, 3) == 8.5
        assert round(section.page_height.inches, 3) == 11.0
        for margin in (section.top_margin, section.bottom_margin,
                       section.left_margin, section.right_margin):
            assert round(margin.inches, 3) == 1.0
        assert round(section.header_distance.inches, 3) == 0.492
        assert round(section.footer_distance.inches, 3) == 0.492

    def test_docx_font_and_heading_tokens(self, tmp_path):
        from docx.shared import RGBColor

        doc = self._docx(tmp_path)
        normal = doc.styles["Normal"]
        assert normal.font.name == "Calibri"
        assert normal.font.size.pt == 11
        assert round(normal.paragraph_format.line_spacing, 2) == 1.10
        assert normal.paragraph_format.space_after.pt == 6

        expected = {
            "Heading 1": (16, "2E74B5", 16, 8),
            "Heading 2": (13, "2E74B5", 12, 6),
            "Heading 3": (12, "1F4D78", 8, 4),
        }
        for name, (size, color, before, after) in expected.items():
            style = doc.styles[name]
            assert style.font.size.pt == size
            assert style.font.bold is True
            assert style.font.color.rgb == RGBColor.from_string(color)
            assert style.paragraph_format.space_before.pt == before
            assert style.paragraph_format.space_after.pt == after

    def test_docx_bullet_numbering_indent(self, tmp_path):
        from docx.oxml.ns import qn

        doc = self._docx(tmp_path)
        num_id = None
        for paragraph in doc.paragraphs:
            ppr = paragraph._p.find(qn("w:pPr"))
            numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
            if numpr is not None:
                num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
                break
        assert num_id is not None, "列表项必须挂真实 Word 编号定义"

        numbering = doc.part.numbering_part.element
        abstract_id = next(
            num.find(qn("w:abstractNumId")).get(qn("w:val"))
            for num in numbering.findall(qn("w:num"))
            if num.get(qn("w:numId")) == num_id
        )
        abstract = next(
            node for node in numbering.findall(qn("w:abstractNum"))
            if node.get(qn("w:abstractNumId")) == abstract_id
        )
        ind = abstract.find(".//" + qn("w:ind"))
        # 标记/文本/悬挂：0.25in / 0.5in / 0.25in，换算成 DXA
        assert ind.get(qn("w:left")) == "720"
        assert ind.get(qn("w:hanging")) == "360"

    def test_docx_table_geometry_is_9360_dxa(self, tmp_path):
        from docx.oxml.ns import qn

        table = self._docx(tmp_path).tables[0]
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "120"
        assert tbl_pr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
        grid = [int(col.get(qn("w:w"))) for col in table._tbl.find(qn("w:tblGrid"))]
        assert sum(grid) == 9360
        first_row = [int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w")))
                     for cell in table.rows[0].cells]
        assert sum(first_row) == 9360

    def test_docx_table_header_and_page_field(self, tmp_path):
        from docx.oxml.ns import qn

        doc = self._docx(tmp_path)
        table = doc.tables[0]
        assert table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "F2F4F7"
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        footer = doc.sections[0].footer.paragraphs[0]._p.xml
        assert 'w:instrText' in footer and "PAGE" in footer

    def test_xlsx_header_and_navigation(self, tmp_path):
        target = tmp_path / "layout.xlsx"
        write_file(target, "excel", parse_markdown(SAMPLE), SAMPLE)
        from openpyxl import load_workbook

        sheet = load_workbook(str(target))["表1"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
        header = sheet["A1"]
        assert header.font.bold is True
        assert header.fill.fgColor.rgb.endswith("F2F4F7")

    def test_pptx_titles_use_real_placeholder(self, tmp_path):
        target = tmp_path / "layout.pptx"
        write_file(target, "ppt", parse_markdown(SAMPLE), SAMPLE)
        from pptx import Presentation

        prs = Presentation(str(target))
        assert round(prs.slide_width / 914400, 3) == 13.333
        assert round(prs.slide_height / 914400, 3) == 7.5
        for slide in prs.slides:
            # Title Only 版式：标题必须落在真实占位符，大纲视图/朗读模式才读得到
            assert slide.shapes.title is not None
            assert slide.shapes.title.text.strip()

    def test_default_word_template_matches_tokens(self, tmp_path):
        from docx import Document

        template = office_writer.template_dir() / "word.docx"
        assert template.exists(), "默认品牌模板应入库"
        doc = Document(str(template))
        round_trip = doc.sections[0]
        assert round(round_trip.page_width.inches, 3) == 8.5
        assert round(round_trip.left_margin.inches, 3) == 1.0
        assert doc.styles["Normal"].font.size.pt == 11

        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        for marker in ("{{ title }}", "{{ subtitle }}", "for s in sections", "s.tables"):
            assert marker in text
