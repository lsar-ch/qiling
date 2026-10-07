"""LocalOfficeAdapter 与办公降级逻辑测试。"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from app.agents.base import AgentTaskRequest
from app.agents.local_office import (
    LocalOfficeAdapter,
    conversion_target,
    _find_office_files,
    _read_docx,
    _read_xlsx,
    _read_pptx,
)


@pytest.fixture
def tmp_office_dir(tmp_path):
    """创建临时办公文件用于测试。"""
    from docx import Document
    doc = Document()
    doc.add_heading("测试标题", level=1)
    doc.add_paragraph("这是第一段正文内容。")
    doc.add_paragraph("这是第二段。")
    doc.save(str(tmp_path / "test.docx"))

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "年龄", "城市"])
    ws.append(["张三", 25, "北京"])
    ws.append(["李四", 30, "上海"])
    ws.append(["王五", 28, "广州"])
    wb.save(str(tmp_path / "test.xlsx"))

    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "首页标题"
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "副标题内容"
    slide2 = prs.slides.add_slide(prs.slide_layouts[0])
    slide2.shapes.title.text = "第二页"
    if len(slide2.placeholders) > 1:
        slide2.placeholders[1].text = "第二页正文"
    prs.save(str(tmp_path / "test.pptx"))

    return tmp_path


class TestFindOfficeFiles:
    def test_finds_docx(self, tmp_office_dir):
        files = _find_office_files(str(tmp_office_dir), "word")
        assert len(files) == 1
        assert files[0].name == "test.docx"

    def test_finds_xlsx(self, tmp_office_dir):
        files = _find_office_files(str(tmp_office_dir), "excel")
        assert len(files) == 1
        assert files[0].name == "test.xlsx"

    def test_finds_pptx(self, tmp_office_dir):
        files = _find_office_files(str(tmp_office_dir), "ppt")
        assert len(files) == 1
        assert files[0].name == "test.pptx"

    def test_returns_empty_for_unknown_kind(self, tmp_office_dir):
        files = _find_office_files(str(tmp_office_dir), "unknown")
        assert files == []

    def test_skips_hidden_dirs(self, tmp_path):
        hidden = tmp_path / ".hidden"
        hidden.mkdir()
        from docx import Document
        doc = Document()
        doc.add_paragraph("hidden")
        doc.save(str(hidden / "secret.docx"))
        files = _find_office_files(str(tmp_path), "word")
        assert files == []


class TestReadDocx:
    def test_reads_heading_and_body(self, tmp_office_dir):
        result = _read_docx(tmp_office_dir / "test.docx")
        assert "测试标题" in result
        assert "第一段正文" in result
        assert "第二段" in result

    def test_empty_doc(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.save(str(tmp_path / "empty.docx"))
        result = _read_docx(tmp_path / "empty.docx")
        assert "空" in result or result == ""


class TestReadXlsx:
    def test_reads_header_and_data(self, tmp_office_dir):
        result = _read_xlsx(tmp_office_dir / "test.xlsx")
        assert "数据" in result
        assert "姓名" in result
        assert "张三" in result
        assert "北京" in result


class TestReadPptx:
    def test_reads_slides(self, tmp_office_dir):
        result = _read_pptx(tmp_office_dir / "test.pptx")
        assert "首页标题" in result
        assert "第二页" in result


class TestLocalOfficeAdapter:
    def test_detect_returns_ready(self):
        adapter = LocalOfficeAdapter()
        info = asyncio.run(adapter.detect())
        assert info.integration_status == "ready"
        assert info.available is True

    def test_run_reads_docx(self, tmp_office_dir):
        adapter = LocalOfficeAdapter()
        events = []

        async def on_update(kind, payload):
            events.append((kind, payload))

        async def run():
            request = AgentTaskRequest(
                task_id="test_001",
                prompt="读取这个文档",
                cwd=str(tmp_office_dir),
                config={"work_kind": "word"},
            )
            return await adapter.run(request, on_update)

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "测试标题" in result.message
        assert len(events) > 0

    def test_run_reads_xlsx(self, tmp_office_dir):
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            request = AgentTaskRequest(
                task_id="test_002",
                prompt="读取表格",
                cwd=str(tmp_office_dir),
                config={"work_kind": "excel"},
            )
            return await adapter.run(request, on_update)

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "张三" in result.message

    def test_run_fails_for_unsupported(self, tmp_office_dir):
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            request = AgentTaskRequest(
                task_id="test_003",
                prompt="读取代码",
                cwd=str(tmp_office_dir),
                config={"work_kind": "code"},
            )
            return await adapter.run(request, on_update)

        result = asyncio.run(run())
        assert result.status == "failed"
        assert "不支持" in result.error

    def test_run_fails_when_no_files(self, tmp_path):
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            request = AgentTaskRequest(
                task_id="test_004",
                prompt="读取文档",
                cwd=str(tmp_path),
                config={"work_kind": "word"},
            )
            return await adapter.run(request, on_update)

        result = asyncio.run(run())
        assert result.status == "failed"
        assert "未找到" in result.error


# ── 2026-10-01 扩充：PDF / 图片 / 邮件 ─────────────────────────────


def _write_min_pdf(path, text: str = "Hello 栖灵 PDF") -> None:
    """手写一份最小可用 PDF（含一段可提取文本），避免为测试引入 reportlab。"""
    stream = f"BT /F1 24 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    Path(path).write_bytes(bytes(out))


def _write_min_png(path, size=(64, 48)) -> None:
    from PIL import Image

    Image.new("RGB", size, (200, 120, 60)).save(str(path))


def _write_min_eml(path) -> None:
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["Subject"] = "周报汇总"
    msg["From"] = "boss@example.com"
    msg["To"] = "me@example.com"
    msg["Date"] = "Wed, 01 Oct 2026 09:00:00 +0800"
    msg.set_content("本周完成了三件事：ACP 接入、办公扩充、回归测试。")
    msg.add_attachment(b"col1,col2\n1,2\n", maintype="text", subtype="csv", filename="data.csv")
    Path(path).write_bytes(bytes(msg))


class TestWidenedReaders:
    def test_reads_pdf_text(self, tmp_path):
        from app.agents.local_office import _read_pdf

        target = tmp_path / "report.pdf"
        _write_min_pdf(target, "Quarterly numbers")
        content = _read_pdf(target)
        assert "Quarterly numbers" in content

    def test_pdf_without_text_hints_ocr(self, tmp_path):
        from PIL import Image
        from app.agents.local_office import _read_pdf

        # 用一张图存成 PDF 近似扫描件；pypdf 提不到文字。
        target = tmp_path / "scan.pdf"
        Image.new("RGB", (32, 32), (10, 10, 10)).save(str(target))
        out = _read_pdf(target)
        assert "PDF" in out or "OCR" in out

    def test_reads_image_metadata(self, tmp_path):
        from app.agents.local_office import _read_image

        target = tmp_path / "pic.png"
        _write_min_png(target, (64, 48))
        content = _read_image(target)
        assert "64" in content and "48" in content
        assert "PNG" in content

    def test_reads_email_headers_body_and_attachments(self, tmp_path):
        from app.agents.local_office import _read_email

        target = tmp_path / "mail.eml"
        _write_min_eml(target)
        content = _read_email(target)
        assert "周报汇总" in content
        assert "boss@example.com" in content
        assert "ACP 接入" in content
        assert "data.csv" in content

    def test_finds_pdf_image_email_files(self, tmp_path):
        _write_min_pdf(tmp_path / "a.pdf")
        _write_min_png(tmp_path / "b.png")
        _write_min_eml(tmp_path / "c.eml")
        assert _find_office_files(str(tmp_path), "pdf")[0].name == "a.pdf"
        assert _find_office_files(str(tmp_path), "image")[0].name == "b.png"
        assert _find_office_files(str(tmp_path), "email")[0].name == "c.eml"

    def test_local_adapter_handles_pdf_work_kind(self, tmp_path):
        adapter = LocalOfficeAdapter()
        _write_min_pdf(tmp_path / "handbook.pdf", "Employee handbook")

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="test_pdf",
                    prompt="总结这份 PDF",
                    cwd=str(tmp_path),
                    config={"work_kind": "pdf"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "Employee handbook" in result.message


class _FakeConvertProc:
    def __init__(self, returncode: int) -> None:
        self.returncode = returncode
        self.stdout = ""
        self.stderr = ""


class TestConversionIntent:
    def test_target_is_read_after_the_convert_verb(self):
        assert conversion_target("把这个 docx 转成 PDF") == "pdf"
        assert conversion_target("把 report.docx 导出为 xlsx") == "excel"
        assert conversion_target("convert notes.docx to pptx") == "ppt"

    def test_source_format_is_not_mistaken_for_target(self):
        # 目标只看「转成」之后，源格式（pdf）不算
        assert conversion_target("把这份 PDF 转成 docx") == "word"

    def test_markdown_target_and_source(self):
        assert conversion_target("把 report.docx 转成 markdown") == "markdown"
        assert conversion_target("把 readme.md 转成 docx") == "word"

    def test_plain_read_has_no_conversion_target(self):
        assert conversion_target("总结这个 docx") is None


class TestLocalOfficeConversion:
    def test_converts_named_file_via_libreoffice(self, tmp_office_dir, monkeypatch):
        from app.agents import office_tools

        def fake_run(command, **kwargs):
            (tmp_office_dir / "test.pdf").write_bytes(b"%PDF-")
            return _FakeConvertProc(0)

        monkeypatch.setattr(
            office_tools.shutil, "which",
            lambda name: "/fake/soffice" if name == "soffice" else None,
        )
        monkeypatch.setattr(office_tools.subprocess, "run", fake_run)

        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="c1",
                    prompt="把 test.docx 转成 PDF",
                    cwd=str(tmp_office_dir),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "test.pdf" in result.message

    def test_conversion_reports_missing_libreoffice(self, tmp_office_dir, monkeypatch):
        from app.agents import office_tools

        monkeypatch.setattr(office_tools.shutil, "which", lambda name: None)
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="c2",
                    prompt="把 test.docx 转成 PDF",
                    cwd=str(tmp_office_dir),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "failed"
        assert "LibreOffice" in result.error


    def test_markdown_to_pdf_uses_builtin_renderer(self, tmp_path, monkeypatch):
        """md -> PDF 不该依赖 Pandoc/LaTeX，两个引擎都没装也能出 PDF。"""
        from app.agents import office_tools

        (tmp_path / "notes.md").write_text("# 标题\n\n- 一\n- 二\n", encoding="utf-8")
        monkeypatch.setattr(office_tools.shutil, "which", lambda name: None)
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="c3",
                    prompt="把 notes.md 转成 PDF",
                    cwd=str(tmp_path),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        assert (tmp_path / "notes.pdf").read_bytes()[:5] == b"%PDF-"

    def test_converts_docx_to_markdown_via_pandoc(self, tmp_office_dir, monkeypatch):
        from app.agents import office_tools

        def fake_run(command, **kwargs):
            (tmp_office_dir / "test.md").write_text("# test\n", encoding="utf-8")
            return _FakeConvertProc(0)

        monkeypatch.setattr(
            office_tools.shutil, "which",
            lambda name: "/fake/pandoc" if name == "pandoc" else None,
        )
        monkeypatch.setattr(office_tools.subprocess, "run", fake_run)
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="c4",
                    prompt="把 test.docx 转成 markdown",
                    cwd=str(tmp_office_dir),
                    config={"work_kind": "word"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "test.md" in result.message


    def test_converts_markdown_to_pptx_via_marp(self, tmp_path, monkeypatch):
        from app.agents import office_tools

        (tmp_path / "deck.md").write_text("# Slide\n\n- a\n", encoding="utf-8")

        def fake_run(command, **kwargs):
            (tmp_path / "deck.pptx").write_bytes(b"PK")
            return _FakeConvertProc(0)

        monkeypatch.setattr(
            office_tools.shutil, "which",
            lambda name: "/fake/marp" if name == "marp" else None,
        )
        monkeypatch.setattr(office_tools.subprocess, "run", fake_run)
        adapter = LocalOfficeAdapter()

        async def on_update(kind, payload):
            pass

        async def run():
            return await adapter.run(
                AgentTaskRequest(
                    task_id="c5",
                    prompt="把 deck.md 转成 pptx",
                    cwd=str(tmp_path),
                    config={"work_kind": "ppt"},
                ),
                on_update,
            )

        result = asyncio.run(run())
        assert result.status == "completed"
        assert "deck.pptx" in result.message
        assert "Marp" in result.message
